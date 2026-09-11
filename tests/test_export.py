"""Synthetic unit-test rows only; never used by submission generation."""
from datetime import datetime, timedelta, timezone
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.utils.exceptions import IllegalCharacterError
from sqlalchemy import func, select

from src.export import TABS, export_xlsx
from src.storage.models import CanonicalEntity, EntityMappingLog, Job, News, Product, RawDocument, ResearchPaper, Startup

NOW = datetime(2026, 9, 11, 12, tzinfo=timezone.utc)


def rows(book, name):
    values = list(book[name].values)
    return [dict(zip(values[0], row)) for row in values[1:]]


@pytest.mark.asyncio
async def test_empty_database_has_exact_six_headers_and_no_padding(db_session, tmp_path):
    report = await export_xlsx(db_session, tmp_path / "empty.xlsx", reference_time=NOW)
    book = load_workbook(tmp_path / "empty.xlsx")
    assert book.sheetnames == [name for name, _ in TABS]
    assert all(sheet.max_row == 1 for sheet in book)
    assert all(v["rows"] == 0 for v in report["tabs"].values())
    assert report["google_sheet_created"] is False
    assert len(report["sha256"]) == 64


@pytest.mark.asyncio
async def test_all_fields_provenance_mapping_and_literal_formulas(db_session, tmp_path):
    entity = CanonicalEntity(canonical_name="Source Co", normalized_name="source co")
    raw = RawDocument(source_name="source", source_url="https://example.org/api", canonical_url="https://example.org/api", content_hash="a" * 64, http_status=200)
    db_session.add_all([entity, raw])
    await db_session.flush()
    common = dict(source_name="source", raw_document_id=raw.id)
    db_session.add_all([
        Startup(entity_name="=HYPERLINK(\"bad\")", source_url="https://example.org/company", canonical_entity_id=entity.id, **common),
        Product(product_name="Source product", startup_name="Source Co", source_url="https://example.org/product", pricing_model=None, **common),
        ResearchPaper(title="Source paper", authors=["Source author"], paper_url="https://example.org/paper", **common),
        Job(company="Source Co", title="ML engineer", url="https://example.org/job", posted_at=NOW, **common),
        News(title="Source news", url="https://example.org/news", published_at=NOW, extracted_metadata={"full_text": "Source body"}, **common),
        EntityMappingLog(raw_name="Source Co", normalized_name="source co", canonical_name="Source Co", canonical_entity_id=entity.id, method="created", confidence=1.0, source_url="https://example.org/company"),
    ])
    await db_session.commit()
    path = tmp_path / "output.xlsx"
    report = await export_xlsx(db_session, path, reference_time=NOW)
    book = load_workbook(path)
    assert all(v["rows"] == 1 for v in report["tabs"].values())
    startup = rows(book, "Startups")[0]
    assert startup["provenance_content_hash"] == "a" * 64
    assert startup["provenance_source_url"] == raw.source_url
    assert startup["resolved_canonical_name"] == "Source Co"
    assert startup["collected_at"].endswith("+00:00")
    assert rows(book, "Products")[0]["pricing_model"] is None
    assert rows(book, "Entity Mapping Log")[0]["source_url"] == "https://example.org/company"
    assert rows(book, "News")[0]["full_text"] == "Source body"
    assert all(cell.data_type != "f" for sheet in book for row in sheet for cell in row)
    assert startup["entity_name"] == '=HYPERLINK("bad")'
    assert await db_session.scalar(select(func.count()).select_from(Startup)) == 1


@pytest.mark.asyncio
@pytest.mark.parametrize("model,date_field,tab", [(Job, "posted_at", "Jobs"), (News, "published_at", "News")])
async def test_export_rechecks_exact_inclusive_24h_and_future_zero(db_session, tmp_path, model, date_field, tab):
    dates = [NOW, NOW - timedelta(hours=24), NOW - timedelta(hours=24, microseconds=1), NOW + timedelta(microseconds=1)]
    for i, date in enumerate(dates):
        fields = dict(title=str(i), url=f"https://example.org/{i}", source_name="source", **{date_field: date})
        if model is Job:
            fields["company"] = "Source company"
        db_session.add(model(**fields))
    await db_session.commit()
    report = await export_xlsx(db_session, tmp_path / "fresh.xlsx", reference_time=NOW)
    exported = rows(load_workbook(tmp_path / "fresh.xlsx"), tab)
    assert {row["title"] for row in exported} == {"0", "1"}
    assert report["tabs"][tab]["excluded_by_freshness"] == 2
    assert report["tabs"][tab]["missing_provenance"] == 2
    assert await db_session.scalar(select(func.count()).select_from(model)) == 4


@pytest.mark.asyncio
async def test_full_text_long_fields_are_lossless_and_paths_confined(db_session, tmp_path):
    raw_root = tmp_path / "raw"
    raw_root.mkdir()
    body = "Real captured content. " * 4000
    (raw_root / "article.txt").write_text(body)
    outside = tmp_path / "private.txt"
    outside.write_text("must not export")
    for i, location in enumerate([raw_root / "article.txt", outside, raw_root / "missing.txt"]):
        db_session.add(News(title=str(i), url=f"https://example.org/{i}", source_name="source", published_at=NOW, full_text_location=str(location)))
    await db_session.commit()
    report = await export_xlsx(db_session, tmp_path / "text.xlsx", reference_time=NOW, raw_root=raw_root)
    data = {r["title"]: r for r in rows(load_workbook(tmp_path / "text.xlsx"), "News")}
    assert data["0"]["full_text"] + data["0"]["full_text__part_2"] + data["0"]["full_text__part_3"] == body
    assert data["1"]["full_text"] is None
    assert data["1"]["full_text_status"] == "outside_raw_root"
    assert data["2"]["full_text_status"] == "unavailable"
    assert report["tabs"]["News"]["missing_full_text"] == 2


@pytest.mark.asyncio
async def test_invalid_text_preserves_existing_artifact(db_session, tmp_path):
    path = tmp_path / "existing.xlsx"
    path.write_bytes(b"previous artifact")
    db_session.add(Startup(entity_name="bad\x00value", source_name="source", source_url="https://example.org"))
    await db_session.commit()
    with pytest.raises(IllegalCharacterError):
        await export_xlsx(db_session, path, reference_time=NOW)
    assert path.read_bytes() == b"previous artifact"


@pytest.mark.asyncio
async def test_rejects_naive_clock_and_wrong_extension(db_session, tmp_path):
    with pytest.raises(ValueError, match="timezone-aware"):
        await export_xlsx(db_session, tmp_path / "x.xlsx", reference_time=NOW.replace(tzinfo=None))
    with pytest.raises(ValueError, match=".xlsx"):
        await export_xlsx(db_session, tmp_path / "x.csv", reference_time=NOW)


@pytest.mark.asyncio
async def test_cli_export_is_wired_and_writes_manifest(db_session, tmp_path, monkeypatch):
    from contextlib import asynccontextmanager
    import json
    from src.main import _run, build_parser

    @asynccontextmanager
    async def scope():
        yield None

    @asynccontextmanager
    async def session_scope():
        yield db_session

    monkeypatch.setattr("src.main.engine_scope", scope)
    monkeypatch.setattr("src.storage.database.get_session_factory", lambda: session_scope)
    path = tmp_path / "cli.xlsx"
    assert await _run(build_parser().parse_args(["--export", "--output", str(path)])) == 0
    assert path.exists()
    manifest = json.loads(path.with_suffix(".manifest.json").read_text())
    assert list(manifest["tabs"]) == [name for name, _ in TABS]
