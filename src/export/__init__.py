"""Read-only, six-tab XLSX export of the existing persistence contract.

No source discovery, enrichment, invented defaults or Google API calls occur here.
The workbook is an assessment-sized snapshot, not a bounded-memory 500k exporter.
"""
from __future__ import annotations

import hashlib
import json
import os
import tempfile
from datetime import datetime, timedelta, timezone
from pathlib import Path
from uuid import UUID

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from src.storage.models import (
    CanonicalEntity, EntityMappingLog, Job, News, Product, RawDocument,
    ResearchPaper, Startup,
)

TABS = (
    ("Startups", Startup), ("Products", Product),
    ("Research Papers", ResearchPaper), ("Jobs", Job),
    ("News", News), ("Entity Mapping Log", EntityMappingLog),
)


def _value(value, *, sqlite: bool):
    if isinstance(value, datetime):
        # SQLite loses tzinfo from DateTime(timezone=True). These repository
        # columns store UTC. This is NOT permission to infer source dates.
        if value.tzinfo is None:
            if not sqlite:
                raise ValueError("Non-SQLite database returned a naive timestamp")
            value = value.replace(tzinfo=timezone.utc)
        return value.astimezone(timezone.utc).isoformat()
    if isinstance(value, UUID):
        return str(value)
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    return value


def _literal_cell(sheet, row: int, column: int, value):
    cell = sheet.cell(row, column, value)
    if isinstance(value, str):
        # Preserve source text exactly but never execute =, +, -, @ as formulas.
        cell.data_type = "s"
    return cell


def _news_text(row: dict, raw_root: Path) -> tuple[str | None, str]:
    inline = (row.get("extracted_metadata") or {}).get("full_text")
    if isinstance(inline, str) and inline:
        return inline, "inline"
    location = row.get("full_text_location")
    if not location:
        return None, "missing"
    path = Path(location).resolve()
    if not path.is_relative_to(raw_root):
        return None, "outside_raw_root"
    try:
        return path.read_text(encoding="utf-8"), "local_file"
    except (OSError, UnicodeError):
        return None, "unavailable"


async def export_xlsx(
    session: AsyncSession,
    output_path: str | Path,
    *,
    reference_time: datetime | None = None,
    raw_root: str | Path = "data/raw",
) -> dict:
    """Export genuine DB rows; recheck News/Jobs at one strict UTC cutoff.

    All domain and mapping fields survive. Provenance/canonical names are LEFT
    JOINed, so missing links remain blank and are counted rather than fabricated.
    Long text is split losslessly into <column>__part_2, etc., not truncated by
    Excel's 32767-character limit. Invalid XML characters fail the export rather
    than silently altering source data. Save atomically; never create DB tables.
    Run against a quiescent DB for a consistent cross-tab assessment snapshot.
    """
    as_of = reference_time if reference_time is not None else datetime.now(timezone.utc)
    if as_of.tzinfo is None or as_of.utcoffset() is None:
        raise ValueError("reference_time must be timezone-aware")
    as_of = as_of.astimezone(timezone.utc)
    output = Path(output_path)
    if output.suffix.lower() != ".xlsx":
        raise ValueError("output_path must end in .xlsx")
    raw_root = Path(raw_root).resolve()
    sqlite = session.get_bind().dialect.name == "sqlite"
    book = Workbook()
    book.remove(book.active)
    book.properties.title = "GraphOne / FrontierAtlas - source-backed snapshot"
    book.properties.description = f"Strict news/jobs window: {as_of - timedelta(hours=24)} through {as_of}; UTC, inclusive."
    report = {
        "output": str(output), "exported_at_utc": as_of.isoformat(),
        "freshness_window_hours": 24, "future_tolerance_seconds": 0,
        "google_sheet_created": False, "tabs": {},
        "text_encoding": "Literal XLSX strings; long fields concatenate base column then __part_2, __part_3, etc.",
    }
    for name, model in TABS:
        columns = list(model.__table__.columns)
        query = select(*columns)
        if model is not EntityMappingLog:
            columns += [c.label(f"provenance_{c.name}") for c in RawDocument.__table__.columns]
            query = select(*columns).outerjoin(RawDocument, model.raw_document_id == RawDocument.id)
        if hasattr(model, "canonical_entity_id"):
            canonical = CanonicalEntity.canonical_name.label("resolved_canonical_name")
            columns.append(canonical)
            query = query.add_columns(canonical).outerjoin(
                CanonicalEntity, model.canonical_entity_id == CanonicalEntity.id
            )
        total = await session.scalar(select(func.count()).select_from(model))
        date_column = Job.posted_at if model is Job else News.published_at if model is News else None
        if date_column is not None:
            query = query.where(date_column >= as_of - timedelta(hours=24), date_column <= as_of)
        query = query.order_by(model.id)
        headers = [c.key for c in columns] + ["export_as_of_utc"]
        if model is News:
            headers += ["full_text", "full_text_status"]
        sheet = book.create_sheet(name)
        sheet.append(headers)
        positions = {key: i + 1 for i, key in enumerate(headers)}
        counts = {"database_rows": total, "rows": 0, "excluded_by_freshness": 0,
                  "missing_provenance": 0, "missing_full_text": 0}
        result = await session.stream(query.execution_options(yield_per=200))
        try:
            async for item in result.mappings():
                data = dict(item)
                data["export_as_of_utc"] = as_of.isoformat()
                if model is not EntityMappingLog and data.get("provenance_id") is None:
                    counts["missing_provenance"] += 1
                if model is News:
                    data["full_text"], data["full_text_status"] = _news_text(data, raw_root)
                    counts["missing_full_text"] += not bool(data["full_text"])
                counts["rows"] += 1
                row_number = counts["rows"] + 1
                if row_number > 1048576:
                    raise ValueError("Excel row limit exceeded; refusing to truncate records")
                for key, value in data.items():
                    value = _value(value, sqlite=sqlite)
                    parts = [value]
                    if isinstance(value, str) and len(value) > 32767:
                        parts = [value[i:i + 32767] for i in range(0, len(value), 32767)]
                    for index, part in enumerate(parts):
                        part_key = key if index == 0 else f"{key}__part_{index + 1}"
                        if part_key not in positions:
                            positions[part_key] = len(positions) + 1
                            _literal_cell(sheet, 1, positions[part_key], part_key)
                        if positions[part_key] > 16384:
                            raise ValueError("Excel column limit exceeded")
                        _literal_cell(sheet, row_number, positions[part_key], part)
        finally:
            await result.close()
        if date_column is not None:
            counts["excluded_by_freshness"] = total - counts["rows"]
        report["tabs"][name] = counts
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="17365D")
            sheet.column_dimensions[cell.column_letter].width = 25
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(dir=output.parent, suffix=".xlsx", delete=False) as temp:
            temporary = Path(temp.name)
        book.save(temporary)
        os.replace(temporary, output)
    finally:
        book.close()
        if temporary is not None:
            temporary.unlink(missing_ok=True)
    report["sha256"] = hashlib.sha256(output.read_bytes()).hexdigest()
    return report
