"""Low-level, read-only reader for the six-tab assessment workbook.

Responsibilities kept deliberately narrow:

* open the workbook read-only and never write to it;
* expose the exact sheet names and their header rows;
* rejoin the exporter's ``<column>__part_N`` continuation columns losslessly
  (the exporter splits any string longer than Excel's 32767-character cell
  limit, see ``src/export/__init__.py``);
* coerce cell text into Python values *strictly* -- an unparseable UUID,
  timestamp or JSON blob raises instead of being silently dropped.

No policy decisions live here. Nothing here touches a database.
"""
from __future__ import annotations

import json
import re
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterator
from uuid import UUID

from openpyxl import load_workbook

# Exact sheet names, in the exact order the exporter emits them.
SHEET_ORDER: tuple[str, ...] = (
    "Startups",
    "Products",
    "Research Papers",
    "Jobs",
    "News",
    "Entity Mapping Log",
)

_PART_RE = re.compile(r"^(?P<base>.+)__part_(?P<index>\d+)$")


class WorkbookFormatError(ValueError):
    """The workbook cannot be read as the expected six-tab export."""


@dataclass(frozen=True)
class Cell:
    """A single logical (post-continuation-join) value and where it came from."""

    sheet: str
    row: int
    column: str
    value: Any

    def where(self) -> str:
        return f"{self.sheet}!{self.column} (row {self.row})"


def _join_parts(sheet: str, row_number: int, header: list[str], values: tuple) -> dict[str, Any]:
    """Collapse ``col``, ``col__part_2``, ``col__part_3`` ... into one string.

    Continuation columns are joined in strict numeric order. A gap (part_3
    present while part_2 is missing) is a corrupted export and raises, because
    silently concatenating what is left would change the source text.
    """
    base: dict[str, Any] = {}
    parts: dict[str, dict[int, str]] = {}

    for name, value in zip(header, values):
        match = _PART_RE.match(name)
        if match is None:
            base[name] = value
            continue
        parts.setdefault(match["base"], {})[int(match["index"])] = value

    for column, indexed in parts.items():
        if column not in base:
            raise WorkbookFormatError(
                f"{sheet} row {row_number}: continuation column '{column}__part_*' has no base column '{column}'"
            )
        head = base[column]
        if head is None:
            # Nothing to continue: every continuation cell must also be empty.
            if any(v not in (None, "") for v in indexed.values()):
                raise WorkbookFormatError(
                    f"{sheet} row {row_number}: '{column}' is empty but its continuation columns are not"
                )
            continue
        chunks = [str(head)]
        for index in sorted(indexed):
            value = indexed[index]
            if value in (None, ""):
                continue
            if index != len(chunks) + 1:
                raise WorkbookFormatError(
                    f"{sheet} row {row_number}: '{column}' continuation parts are not contiguous "
                    f"(expected __part_{len(chunks) + 1}, found __part_{index})"
                )
            chunks.append(str(value))
        base[column] = "".join(chunks)

    return base


def read_sheet(path: str | Path, sheet: str) -> tuple[list[str], list[dict[str, Any]]]:
    """Return ``(header, rows)`` for one sheet, continuation columns rejoined."""
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet not in book.sheetnames:
            raise WorkbookFormatError(f"missing sheet '{sheet}'")
        worksheet = book[sheet]
        iterator: Iterator[tuple] = worksheet.iter_rows(values_only=True)
        try:
            raw_header = next(iterator)
        except StopIteration as exc:  # pragma: no cover - an empty sheet has no header row
            raise WorkbookFormatError(f"sheet '{sheet}' has no header row") from exc
        header = [str(h) for h in raw_header if h is not None]
        if len(header) != len(set(header)):
            raise WorkbookFormatError(f"sheet '{sheet}' has duplicate headers")
        rows: list[dict[str, Any]] = []
        for offset, values in enumerate(iterator, start=2):
            if all(v is None for v in values):
                continue  # trailing blank row left behind by a spreadsheet editor
            rows.append(_join_parts(sheet, offset, header, values))
        return header, rows
    finally:
        book.close()


def read_workbook(path: str | Path) -> dict[str, tuple[list[str], list[dict[str, Any]]]]:
    """Read all six sheets. Raises if a sheet is missing or extra."""
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        found = tuple(book.sheetnames)
    finally:
        book.close()
    if found != SHEET_ORDER:
        raise WorkbookFormatError(
            "unexpected sheets: expected exactly "
            f"{list(SHEET_ORDER)} in order, found {list(found)}"
        )
    return {name: read_sheet(path, name) for name in SHEET_ORDER}


# ---------------------------------------------------------------------------
# Strict cell coercion
# ---------------------------------------------------------------------------


def blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and value.strip() == "")


def as_text(cell: Cell, *, required: bool = False, max_length: int | None = None) -> str | None:
    if blank(cell.value):
        if required:
            raise WorkbookFormatError(f"{cell.where()}: required text value is empty")
        return None
    text = cell.value if isinstance(cell.value, str) else str(cell.value)
    if max_length is not None and len(text) > max_length:
        raise WorkbookFormatError(
            f"{cell.where()}: value is {len(text)} characters, exceeding the column limit of {max_length}"
        )
    return text


def as_uuid(cell: Cell, *, required: bool = False) -> UUID | None:
    if blank(cell.value):
        if required:
            raise WorkbookFormatError(f"{cell.where()}: required UUID is empty")
        return None
    if isinstance(cell.value, UUID):
        return cell.value
    try:
        return UUID(str(cell.value))
    except (ValueError, AttributeError, TypeError) as exc:
        raise WorkbookFormatError(f"{cell.where()}: {cell.value!r} is not a valid UUID") from exc


def as_datetime(cell: Cell, *, required: bool = False) -> datetime | None:
    """Parse an ISO-8601 timestamp and require it to be timezone-aware."""
    if blank(cell.value):
        if required:
            raise WorkbookFormatError(f"{cell.where()}: required timestamp is empty")
        return None
    value = cell.value
    if isinstance(value, datetime):
        parsed = value
    else:
        try:
            parsed = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
        except ValueError as exc:
            raise WorkbookFormatError(f"{cell.where()}: {value!r} is not an ISO-8601 timestamp") from exc
    if parsed.tzinfo is None or parsed.utcoffset() is None:
        raise WorkbookFormatError(
            f"{cell.where()}: timestamp {value!r} is timezone-naive; the export contract is timezone-aware UTC"
        )
    return parsed.astimezone(timezone.utc)


def as_int(cell: Cell, *, required: bool = False, minimum: int | None = None) -> int | None:
    if blank(cell.value):
        if required:
            raise WorkbookFormatError(f"{cell.where()}: required integer is empty")
        return None
    value = cell.value
    if isinstance(value, bool):
        raise WorkbookFormatError(f"{cell.where()}: {value!r} is a boolean, not an integer")
    if isinstance(value, int):
        number = value
    elif isinstance(value, float) and value.is_integer():
        number = int(value)
    else:
        try:
            number = int(str(value).strip())
        except ValueError as exc:
            raise WorkbookFormatError(f"{cell.where()}: {value!r} is not an integer") from exc
    if minimum is not None and number < minimum:
        raise WorkbookFormatError(f"{cell.where()}: {number} is below the allowed minimum {minimum}")
    return number


def as_float(cell: Cell, *, required: bool = False) -> float | None:
    if blank(cell.value):
        if required:
            raise WorkbookFormatError(f"{cell.where()}: required number is empty")
        return None
    try:
        return float(cell.value)
    except (TypeError, ValueError) as exc:
        raise WorkbookFormatError(f"{cell.where()}: {cell.value!r} is not a number") from exc


def as_bool(cell: Cell, *, required: bool = False) -> bool | None:
    if blank(cell.value):
        if required:
            raise WorkbookFormatError(f"{cell.where()}: required boolean is empty")
        return None
    value = cell.value
    if isinstance(value, bool):
        return value
    text = str(value).strip().lower()
    if text in {"true", "1"}:
        return True
    if text in {"false", "0"}:
        return False
    raise WorkbookFormatError(f"{cell.where()}: {value!r} is not a boolean")


def as_json(cell: Cell, *, expected: type | tuple[type, ...], default: Any) -> Any:
    """Parse a JSON cell, requiring the documented container shape.

    The exporter writes ``json.dumps`` of a dict or list. Anything else (a bare
    scalar, or malformed JSON) is a corrupted cell and raises -- it is never
    coerced into an empty container, because that would silently lose data.
    """
    if blank(cell.value):
        return default
    try:
        parsed = json.loads(cell.value)
    except (TypeError, json.JSONDecodeError) as exc:
        raise WorkbookFormatError(f"{cell.where()}: value is not parseable JSON ({exc})") from exc
    if not isinstance(parsed, expected):
        names = expected.__name__ if isinstance(expected, type) else "/".join(t.__name__ for t in expected)
        raise WorkbookFormatError(
            f"{cell.where()}: JSON parsed as {type(parsed).__name__}, expected {names}"
        )
    return parsed
